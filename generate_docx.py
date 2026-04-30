from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─── Color Palette ───────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x0D, 0x2B, 0x55)
MID_BLUE    = RGBColor(0x1A, 0x5F, 0x9E)
ACCENT_BLUE = RGBColor(0x26, 0x8B, 0xD2)
TABLE_HEAD  = RGBColor(0x1A, 0x5F, 0x9E)
TABLE_ALT   = RGBColor(0xEA, 0xF4, 0xFB)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK   = RGBColor(0x1A, 0x1A, 0x2E)
GRAY_TEXT   = RGBColor(0x55, 0x55, 0x55)

def rgb_hex(color):
    """Convert RGBColor (tuple) to hex string."""
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"

# ─── Helpers ──────────────────────────────────────────────────────────────────
def set_cell_bg(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), rgb_hex(color))
    tcPr.append(shd)

def set_cell_borders(cell, border_color="268BD2"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), border_color)
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_para_shading(p, color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), rgb_hex(color))
    pPr.append(shd)

def add_para_bottom_border(p, color_hex, sz='12'):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), sz)
    bot.set(qn('w:color'), color_hex)
    pBdr.append(bot)
    pPr.append(pBdr)

def add_heading(doc, text, color=DARK_BLUE, size=15, bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    return p

def add_section_banner(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    add_para_shading(p, ACCENT_BLUE)
    run = p.add_run(f"  {text}  ")
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = WHITE
    run.font.name = 'Calibri'
    return p

def add_body(doc, text, size=10.5, color=TEXT_DARK, space_before=2, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = 'Calibri'
    return p

def add_bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = TEXT_DARK
    run.font.name = 'Calibri'
    return p

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_bg(cell, TABLE_HEAD)
        set_cell_borders(cell, "1A5F9E")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = WHITE
        run.font.name = 'Calibri'

    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        bg = TABLE_ALT if ri % 2 == 0 else WHITE
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            set_cell_bg(cell, bg)
            set_cell_borders(cell, "AECDE6")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(3)
            run = p.add_run(str(val))
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT_DARK
            run.font.name = 'Calibri'

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    return table

def add_page_number_to_footer(section):
    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_left = fp.add_run("Medi Assist Healthcare Ltd.  |  Opus - Raksha Prime  |  Milestone 2  |  Page ")
    run_left.font.size = Pt(8.5)
    run_left.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    run_left.font.name = 'Calibri'

    # Page number field
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run_pg = fp.add_run()
    run_pg.font.size = Pt(8.5)
    run_pg.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    run_pg.font.name = 'Calibri'
    run_pg._r.append(fldChar1)
    run_pg._r.append(instrText)
    run_pg._r.append(fldChar2)


# ═════════════════════════════════════════════════════════════════════════════
# BUILD DOCUMENT
# ═════════════════════════════════════════════════════════════════════════════
doc = Document()

# Page margins
sec = doc.sections[0]
sec.top_margin    = Cm(2.0)
sec.bottom_margin = Cm(1.8)
sec.left_margin   = Cm(2.5)
sec.right_margin  = Cm(2.5)

# ── COVER PAGE ────────────────────────────────────────────────────────────────
# Top navy bar
p_bar = doc.add_paragraph()
p_bar.paragraph_format.space_before = Pt(0)
p_bar.paragraph_format.space_after  = Pt(0)
add_para_shading(p_bar, DARK_BLUE)
run = p_bar.add_run("  ")
run.font.size = Pt(24)

doc.add_paragraph().paragraph_format.space_after = Pt(30)

# Main title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(20)
p_title.paragraph_format.space_after  = Pt(8)
r = p_title.add_run(
    "Performance Evaluation and Optimization Analysis\n"
    "for Raksha Prime in Health Insurance"
)
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = DARK_BLUE
r.font.name = 'Calibri'

# Subtitle badge
p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_before = Pt(4)
p_sub.paragraph_format.space_after  = Pt(28)
r2 = p_sub.add_run("  MILESTONE 2 REPORT  ")
r2.bold = True
r2.font.size = Pt(13)
r2.font.color.rgb = WHITE
r2.font.name = 'Calibri'
add_para_shading(p_sub, ACCENT_BLUE)

# Horizontal rule
p_rule = doc.add_paragraph()
p_rule.paragraph_format.space_before = Pt(0)
p_rule.paragraph_format.space_after  = Pt(20)
add_para_bottom_border(p_rule, "1A5F9E", sz='16')

# Student details table
info = [
    ("Name",           "[Your Name]"),
    ("Roll Number",    "[Your Roll Number]"),
    ("Company",        "Medi Assist Healthcare Ltd."),
    ("Department",     "Opus - Raksha Prime"),
    ("Academic Year",  "[Academic Year]"),
    ("Submitted To",   "Faculty Mentor"),
]
detail_table = doc.add_table(rows=len(info), cols=2)
detail_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (lbl, val) in enumerate(info):
    lc = detail_table.rows[i].cells[0]
    lp = lc.paragraphs[0]
    lp.paragraph_format.space_before = Pt(4)
    lp.paragraph_format.space_after  = Pt(4)
    lr = lp.add_run(lbl)
    lr.bold = True
    lr.font.size = Pt(11)
    lr.font.color.rgb = MID_BLUE
    lr.font.name = 'Calibri'
    lc.width = Inches(2.0)

    vc = detail_table.rows[i].cells[1]
    vp = vc.paragraphs[0]
    vp.paragraph_format.space_before = Pt(4)
    vp.paragraph_format.space_after  = Pt(4)
    vr = vp.add_run(val)
    vr.font.size = Pt(11)
    vr.font.color.rgb = TEXT_DARK
    vr.font.name = 'Calibri'

doc.add_page_break()

# ── SECTION 1: Project Context ────────────────────────────────────────────────
add_heading(doc, "1.  Project Context", size=15)

add_body(doc,
    "Medi Assist Healthcare Ltd. is India's largest healthcare insurance Third Party Administrator (TPA), "
    "facilitating seamless coordination between insurers, hospitals, and policyholders. The organization "
    "manages both cashless and reimbursement claims and leverages technology to enhance efficiency and "
    "reduce turnaround time. The Raksha Prime product focuses on improving the post-discharge experience "
    "and overall operational efficiency.")

add_body(doc,
    "As a Data Analytics Intern in the Opus Department, my role involves analyzing claims and operational "
    "data, processing daily leads, updating Out-of-Pocket (OOP) data, and working with OPUS-tagged "
    "datasets to support performance evaluation.")

add_heading(doc, "Project Objective", color=MID_BLUE, size=11, space_before=6, space_after=2)
add_body(doc,
    "To analyze Raksha Prime operational data to identify inefficiencies, delays, and cost patterns, "
    "and provide data-driven recommendations for process optimization.")

# ── SECTION 2: Key Business Questions ────────────────────────────────────────
add_heading(doc, "2.  Key Business Questions", size=15)
add_body(doc, "This project is designed to answer the following critical questions:")
for bq in [
    "What factors contribute to delays in patient discharge under Raksha Prime?",
    "Which hospitals show abnormal claim or billing patterns?",
    "How does Out-of-Pocket (OOP) vary across cases and hospitals?",
    "Are there inefficiencies in the claims processing workflow?",
    "Can patterns in data indicate potential risks or inconsistencies?",
]:
    add_bullet(doc, bq)

# ── FOCUS AREA 1 ──────────────────────────────────────────────────────────────
add_section_banner(doc, "FOCUS AREA 1: Clarity of Thought & Timeline Management")

add_heading(doc, "3.  Deliverables Mapping", size=14)
add_body(doc, "Each deliverable is explicitly linked to a business question with a defined expected outcome:")
doc.add_paragraph().paragraph_format.space_after = Pt(4)

add_styled_table(doc,
    headers=["Deliverable", "Business Question Addressed", "Expected Outcome"],
    rows=[
        ["Claims Data Analysis",    "Identify delay factors",    "Key drivers of delays"],
        ["OOP Analysis",            "Understand cost variation", "Patient cost insights"],
        ["Hospital-Level Analysis", "Detect abnormal patterns",  "High-risk hospital identification"],
        ["Process Flow Analysis",   "Identify inefficiencies",   "Process improvement areas"],
        ["Final Report",            "Consolidated insights",     "Actionable recommendations"],
    ],
    col_widths=[2.1, 2.5, 2.3]
)

add_heading(doc, "4.  Work Breakdown & Timeline", size=14)
add_styled_table(doc,
    headers=["Phase", "Tasks", "Timeline", "Status"],
    rows=[
        ["Data Understanding",   "Learn claims, OOP, OPUS data",           "Week 1-2", "Completed"],
        ["Data Processing",      "Daily leads, OOP updates, OPUS handling", "Week 2-3", "Ongoing"],
        ["Data Cleaning",        "Handle missing/inconsistent data",        "Week 3-4", "Ongoing"],
        ["Exploratory Analysis", "Identify trends & patterns",              "Week 4-6", "Planned"],
        ["Insight Generation",   "Develop findings",                        "Week 6-7", "Planned"],
        ["Final Report",         "Recommendations & summary",               "Week 7-8", "Planned"],
    ],
    col_widths=[1.7, 2.7, 1.3, 1.2]
)

# ── FOCUS AREA 2 ──────────────────────────────────────────────────────────────
add_section_banner(doc, "FOCUS AREA 2: Content Knowledge")

add_heading(doc, "5.  Data & Information Requirements", size=14)
add_body(doc, "The following data is required to execute the analysis effectively:")
for b in [
    "Claims data — amount, approval status, processing time",
    "Out-of-Pocket (OOP) data",
    "Hospital-level data — tiers, locations, historical performance",
    "Patient-level data — demographics and policy details",
    "Timeline data — admission to discharge markers",
    "OPUS-tagged operational data",
]:
    add_bullet(doc, b)

add_heading(doc, "6.  Data Availability & Extraction Status", size=14)
add_styled_table(doc,
    headers=["Data Type", "Availability", "Source", "Status"],
    rows=[
        ["Claims Data",   "Available", "Internal systems", "Accessed"],
        ["OOP Data",      "Available", "Daily updates",    "Regularly processed"],
        ["Hospital Data", "Available", "OPUS system",      "Partially explored"],
        ["Patient Data",  "Available", "Claims datasets",  "Accessed"],
        ["Timeline Data", "Limited",   "Claims logs",      "Needs structuring"],
    ],
    col_widths=[1.6, 1.4, 1.9, 2.0]
)

add_heading(doc, "7.  Preliminary Insights & Observations", size=14)
add_body(doc, "Based on initial analysis and operational exposure, the following patterns have been observed:")
for ins in [
    "Certain hospitals consistently show higher claim amounts compared to others.",
    "Significant variation exists in Out-of-Pocket (OOP) expenses across similar cases.",
    "Delays in discharge are often linked to documentation gaps and processing inefficiencies.",
    "OPUS-tagged data indicates inconsistencies in handling specific cases.",
    "Initial observations suggest scope for improving turnaround time and process standardization.",
]:
    add_bullet(doc, ins)

# ── FOCUS AREA 3 ──────────────────────────────────────────────────────────────
add_section_banner(doc, "FOCUS AREA 3: Risk Identification")

add_heading(doc, "8.  Risks & Mitigation Plan", size=14)
add_styled_table(doc,
    headers=["Risk", "Impact", "Mitigation Strategy"],
    rows=[
        ["Incomplete or missing data", "Inaccurate insights",       "Data validation and cleaning"],
        ["Data inconsistency",         "Incorrect analysis",        "Standardization of datasets"],
        ["Limited data access",        "Delays in progress",        "Coordination with internal teams"],
        ["Time constraints",           "Incomplete deliverables",   "Strict adherence to timeline"],
        ["Operational complexity",     "Misinterpretation of data", "Continuous consultation with team"],
    ],
    col_widths=[2.0, 2.0, 2.9]
)

add_heading(doc, "9.  Next Steps", size=14)
for ns in [
    "Complete data cleaning and structuring.",
    "Perform exploratory data analysis (EDA).",
    "Identify patterns and key inefficiencies.",
    "Generate actionable insights and recommendations.",
    "Prepare final internship report.",
]:
    add_bullet(doc, ns)

# ── APPENDIX ────────────────────────────────────────────────────────────────
doc.add_page_break()
add_heading(doc, "Appendix", size=15)
add_heading(doc, "Appendix A: Milestone 1 Submission", color=MID_BLUE, size=12, space_before=4, space_after=2)
add_body(doc,
    "(Attach previously submitted Milestone 1 document here for reference, as suggested for continuity.)",
    color=GRAY_TEXT)

# ── FOOTER ───────────────────────────────────────────────────────────────────
for section in doc.sections:
    add_page_number_to_footer(section)

# ── SAVE ────────────────────────────────────────────────────────────────────
output_path = r"c:\Users\Hassain\.gemini\antigravity\scratch\raksha-prime\Milestone2_Report_Final.docx"
doc.save(output_path)
print(f"Saved: {output_path}")
